"""
Test suite for table rendering (markdown tables).

Enforces:
- Invariant 1: Text Fidelity - all table cell content appears in output
- Invariant 2: Determinism - same table structure produces same output
- Invariant 3: Observability - all artifacts written to disk
"""
import pytest
from pathlib import Path
from docx import Document

from docx_template_export.models.export_models import WordExportRequest
from tests.test_artifacts_helpers import (
    run_export_and_collect_artifacts,
    assert_text_fidelity,
    assert_no_placeholders_remain,
)


@pytest.fixture
def artifacts_dir(test_output_dir):
    """Return artifacts directory for this test suite."""
    artifacts_base = test_output_dir.parent / "artifacts"
    artifacts_base.mkdir(exist_ok=True)
    return artifacts_base


@pytest.fixture
def table_template(test_fixtures_dir):
    """Create template with block placeholder for tables."""
    template_path = test_fixtures_dir / "table_template.docx"
    doc = Document()
    doc.add_paragraph("{{table_content}}")
    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(template_path))
    return template_path


def test_markdown_table_simple(artifacts_dir, table_template):
    """Test simple markdown table rendering."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "table_content": """| Column A | Column B | Column C |
|----------|----------|----------|
| Value 1  | Value 2  | Value 3  |
| Value 4  | Value 5  | Value 6  |
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="markdown_table_simple",
        template_path=table_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="tables",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "Column A",
            "Column B",
            "Column C",
            "Value 1",
            "Value 2",
            "Value 3",
            "Value 4",
            "Value 5",
            "Value 6",
        ],
        extracted_text=extracted_text,
        test_name="markdown_table_simple",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["table_content"],
        test_name="markdown_table_simple",
    )


def test_markdown_table_with_text_before_after(artifacts_dir, table_template):
    """Test table with text before and after."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "table_content": """Text before table.

| A | B |
|---|---|
| 1 | 2 |

Text after table.
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="markdown_table_with_text_before_after",
        template_path=table_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="tables",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "Text before table",
            "A",
            "B",
            "1",
            "2",
            "Text after table",
        ],
        extracted_text=extracted_text,
        test_name="markdown_table_with_text_before_after",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["table_content"],
        test_name="markdown_table_with_text_before_after",
    )


def test_table_determinism(artifacts_dir, table_template):
    """Test that same table structure produces identical output (Invariant 2: Determinism)."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "table_content": """| X | Y |
|---|---|
| a | b |
""",
        },
    )
    
    # Run export twice
    artifact_dir1, markdown1, extracted_text1 = run_export_and_collect_artifacts(
        test_name="table_determinism_run1",
        template_path=table_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="tables",
    )
    
    artifact_dir2, markdown2, extracted_text2 = run_export_and_collect_artifacts(
        test_name="table_determinism_run2",
        template_path=table_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="tables",
    )
    
    # Assert markdown is identical
    assert markdown1 == markdown2, "Markdown output must be identical for same input"
    
    # Assert extracted text is identical (normalized)
    normalized1 = " ".join(extracted_text1.split())
    normalized2 = " ".join(extracted_text2.split())
    assert normalized1 == normalized2, "Extracted text must be identical for same input"
