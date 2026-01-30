"""
Test suite for list rendering (bullet and numbered lists).

Enforces:
- Invariant 1: Text Fidelity - all list item text appears in output
- Invariant 2: Determinism - same list structure produces same output
- Invariant 3: Observability - all artifacts written to disk
"""
import pytest
from pathlib import Path
from docx import Document

from docx_template_export.models.export_models import WordExportRequest
from tests.test_artifacts_helpers import (
    run_export_and_collect_artifacts,
    assert_text_fidelity,
    assert_no_placeholders_remain,
)


@pytest.fixture
def artifacts_dir(test_output_dir):
    """Return artifacts directory for this test suite."""
    artifacts_base = test_output_dir.parent / "artifacts"
    artifacts_base.mkdir(exist_ok=True)
    return artifacts_base


@pytest.fixture
def list_template(test_fixtures_dir):
    """Create template with block placeholder for lists."""
    template_path = test_fixtures_dir / "list_template.docx"
    doc = Document()
    doc.add_paragraph("{{list_content}}")
    template_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(template_path))
    return template_path


def test_bullet_list_flat(artifacts_dir, list_template):
    """Test flat bullet list rendering."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- First bullet item
- Second bullet item
- Third bullet item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="bullet_list_flat",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "First bullet item",
            "Second bullet item",
            "Third bullet item",
        ],
        extracted_text=extracted_text,
        test_name="bullet_list_flat",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="bullet_list_flat",
    )


def test_numbered_list_flat(artifacts_dir, list_template):
    """Test flat numbered list rendering."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """1. First numbered item
2. Second numbered item
3. Third numbered item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="numbered_list_flat",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "First numbered item",
            "Second numbered item",
            "Third numbered item",
        ],
        extracted_text=extracted_text,
        test_name="numbered_list_flat",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="numbered_list_flat",
    )


def test_nested_bullet_list(artifacts_dir, list_template):
    """Test nested bullet list rendering."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- Level 1 item
  - Level 2 item
    - Level 3 item
  - Another level 2 item
- Another level 1 item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="nested_bullet_list",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "Level 1 item",
            "Level 2 item",
            "Level 3 item",
            "Another level 2 item",
            "Another level 1 item",
        ],
        extracted_text=extracted_text,
        test_name="nested_bullet_list",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="nested_bullet_list",
    )


def test_nested_numbered_list(artifacts_dir, list_template):
    """Test nested numbered list rendering."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """1. Level 1 item
   1. Level 2 item
      1. Level 3 item
   2. Another level 2 item
2. Another level 1 item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="nested_numbered_list",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "Level 1 item",
            "Level 2 item",
            "Level 3 item",
            "Another level 2 item",
            "Another level 1 item",
        ],
        extracted_text=extracted_text,
        test_name="nested_numbered_list",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="nested_numbered_list",
    )


def test_mixed_bullet_and_numbered_lists(artifacts_dir, list_template):
    """Test that bullet and numbered lists are rendered as separate blocks."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- Bullet item one
- Bullet item two

1. Numbered item one
2. Numbered item two
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="mixed_bullet_and_numbered_lists",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "Bullet item one",
            "Bullet item two",
            "Numbered item one",
            "Numbered item two",
        ],
        extracted_text=extracted_text,
        test_name="mixed_bullet_and_numbered_lists",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="mixed_bullet_and_numbered_lists",
    )


def test_list_with_continuation_paragraphs(artifacts_dir, list_template):
    """Test list items with continuation paragraphs."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- First item with continuation

  This is a continuation paragraph that must not be lost.

- Second item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="list_with_continuation_paragraphs",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity
    assert_text_fidelity(
        input_texts=[
            "First item with continuation",
            "This is a continuation paragraph that must not be lost",
            "Second item",
        ],
        extracted_text=extracted_text,
        test_name="list_with_continuation_paragraphs",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="list_with_continuation_paragraphs",
    )


def test_list_item_with_table(artifacts_dir, list_template):
    """Test list item containing a markdown table as continuation content."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- Item intro paragraph

  | A | B |
  |---|---|
  | 1 | 2 |

  Another continuation paragraph

- Next item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="list_item_with_table",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity - all content must appear
    assert_text_fidelity(
        input_texts=[
            "Item intro paragraph",
            "A",
            "B",
            "1",
            "2",
            "Another continuation paragraph",
            "Next item",
        ],
        extracted_text=extracted_text,
        test_name="list_item_with_table",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="list_item_with_table",
    )
    
    # Verify table content appears (text-grid format)
    assert "A" in extracted_text
    assert "B" in extracted_text
    assert "1" in extracted_text
    assert "2" in extracted_text


def test_numbered_list_item_with_table(artifacts_dir, list_template):
    """Test numbered list item containing a markdown table as continuation content."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """1. Item intro paragraph

   | X | Y |
   |---|---|
   | a | b |

   Another continuation paragraph

2. Next item
""",
        },
    )
    
    artifact_dir, markdown, extracted_text = run_export_and_collect_artifacts(
        test_name="numbered_list_item_with_table",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert text fidelity - all content must appear
    assert_text_fidelity(
        input_texts=[
            "Item intro paragraph",
            "X",
            "Y",
            "a",
            "b",
            "Another continuation paragraph",
            "Next item",
        ],
        extracted_text=extracted_text,
        test_name="numbered_list_item_with_table",
    )
    
    # Assert no placeholders remain
    assert_no_placeholders_remain(
        extracted_text=extracted_text,
        scalar_placeholders=[],
        block_placeholders=["list_content"],
        test_name="numbered_list_item_with_table",
    )
    
    # Verify table content appears (text-grid format)
    assert "X" in extracted_text
    assert "Y" in extracted_text
    assert "a" in extracted_text
    assert "b" in extracted_text


def test_list_determinism(artifacts_dir, list_template):
    """Test that same list structure produces identical output (Invariant 2: Determinism)."""
    request = WordExportRequest(
        scalar_fields={},
        block_fields={
            "list_content": """- Item A
- Item B
- Item C
""",
        },
    )
    
    # Run export twice
    artifact_dir1, markdown1, extracted_text1 = run_export_and_collect_artifacts(
        test_name="list_determinism_run1",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    artifact_dir2, markdown2, extracted_text2 = run_export_and_collect_artifacts(
        test_name="list_determinism_run2",
        template_path=list_template,
        request=request,
        artifacts_dir=artifacts_dir,
        subdirectory="lists",
    )
    
    # Assert markdown is identical
    assert markdown1 == markdown2, "Markdown output must be identical for same input"
    
    # Assert extracted text is identical (normalized)
    normalized1 = " ".join(extracted_text1.split())
    normalized2 = " ".join(extracted_text2.split())
    assert normalized1 == normalized2, "Extracted text must be identical for same input"
